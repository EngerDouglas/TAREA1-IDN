"""Genera el documento de la solución de la Práctica 1 de ICC-321 (entregable 1).

Escribe el .docx con python-docx, lo guarda de nuevo con LibreOffice, lo convierte
a PDF, lee en qué página quedó cada título y vuelve a armar el índice con esos
números. Las figuras del esquema y la arquitectura salen de figuras/*.html.

Requisitos: pip install python-docx; LibreOffice (soffice) y poppler (pdftotext).
Uso:    python3 generar_documento_solucion.py
Salida: 3_Informes/Practica1_Documento_Solucion.docx y .pdf
"""
import os
import re
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(BASE)
FIG = os.path.join(BASE, "figuras")
SALIDA = os.path.join(REPO, "3_Informes")
NOMBRE = "Practica1_Documento_Solucion"

FUENTE = "Times New Roman"


# ------------------------------------------------------------------ utilidades ----

def fuente(run, tam=None, negrita=None, cursiva=None, nombre=FUENTE):
    run.font.name = nombre
    run._element.rPr.rFonts.set(qn("w:eastAsia"), nombre)
    if tam:
        run.font.size = Pt(tam)
    if negrita is not None:
        run.bold = negrita
    if cursiva is not None:
        run.italic = cursiva


def texto_con_formato(par, texto, tam=12):
    """Escribe texto admitiendo **negrita** y *cursiva*."""
    for trozo in re.split(r"(\*\*.+?\*\*|\*.+?\*)", texto):
        if not trozo:
            continue
        if trozo.startswith("**"):
            fuente(par.add_run(trozo[2:-2]), tam, negrita=True)
        elif trozo.startswith("*"):
            fuente(par.add_run(trozo[1:-1]), tam, cursiva=True)
        else:
            fuente(par.add_run(trozo), tam)


class Informe:
    def __init__(self, paginas_indice=None):
        self.doc = Document()
        self.paginas = paginas_indice or {}
        self.titulos = []
        self.nfig = 0
        self.ntab = 0
        self._estilos()

    def _estilos(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
        for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, lado, Cm(2.54))
        normal = d.styles["Normal"]
        normal.font.name = FUENTE
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
        normal.font.size = Pt(12)
        idioma = OxmlElement("w:lang")
        idioma.set(qn("w:val"), "es-DO")
        normal.element.get_or_add_rPr().append(idioma)
        pf = normal.paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(6)
        for nombre, tam in (("Heading 1", 14), ("Heading 2", 12)):
            st = d.styles[nombre]
            st.font.name = FUENTE
            st.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
            st.font.size = Pt(tam)
            st.font.bold = True
            st.font.italic = False
            st.font.color.rgb = RGBColor(0, 0, 0)
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                st.element.rPr.rFonts.attrib.pop(qn(attr), None)
            st.paragraph_format.space_before = Pt(18 if nombre == "Heading 1" else 12)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.line_spacing = 1.15

    # -- bloques --
    def centrado(self, texto, tam=12, negrita=False, antes=0, despues=0):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(antes)
        p.paragraph_format.space_after = Pt(despues)
        p.paragraph_format.line_spacing = 1.15
        fuente(p.add_run(texto), tam, negrita=negrita)
        return p

    def salto(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h1(self, texto):
        self.titulos.append((1, texto))
        self.doc.add_heading(texto, level=1)

    def h2(self, texto):
        self.titulos.append((2, texto))
        self.doc.add_heading(texto, level=2)

    def p(self, texto):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        texto_con_formato(par, texto)
        return par

    def lista(self, items, numerada=False):
        for i, t in enumerate(items, 1):
            par = self.doc.add_paragraph(style="List Number" if numerada else "List Bullet")
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.space_after = Pt(3)
            texto_con_formato(par, t)

    def tabla(self, titulo, encabezado, filas, anchos):
        """Tabla con su título arriba (Tabla N. ...), como se acostumbra."""
        self.ntab += 1
        cap = self.doc.add_paragraph()
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(3)
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.line_spacing = 1.0
        fuente(cap.add_run("Tabla %d. " % self.ntab), 10, negrita=True)
        fuente(cap.add_run(titulo), 10)
        t = self.doc.add_table(rows=1, cols=len(encabezado))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(encabezado):
            c = t.rows[0].cells[i]
            c.text = ""
            fuente(c.paragraphs[0].add_run(h), 10, negrita=True)
            sombra = OxmlElement("w:shd")
            sombra.set(qn("w:val"), "clear")
            sombra.set(qn("w:fill"), "D9D9D9")
            c._tc.get_or_add_tcPr().append(sombra)
        for f in filas:
            celdas = t.add_row().cells
            for i, v in enumerate(f):
                celdas[i].text = ""
                par = celdas[i].paragraphs[0]
                texto_con_formato(par, str(v), 10)
        t.autofit = False
        # LibreOffice toma el ancho de la rejilla y de la tabla, no solo el de cada celda
        for i, w in enumerate(anchos):
            t.columns[i].width = Cm(w)
        tblPr = t._tbl.tblPr
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:type"), "dxa")
        tblW.set(qn("w:w"), str(int(sum(anchos) / 2.54 * 1440)))
        for fila in t.rows:
            for i, w in enumerate(anchos):
                fila.cells[i].width = Cm(w)
            for c in fila.cells:
                for par in c.paragraphs:
                    par.paragraph_format.line_spacing = 1.0
                    par.paragraph_format.space_after = Pt(2)
        # la fila de encabezado se repite si la tabla cambia de página
        trPr = t.rows[0]._tr.get_or_add_trPr()
        rep = OxmlElement("w:tblHeader")
        rep.set(qn("w:val"), "true")
        trPr.append(rep)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def figura(self, archivo, pie, ancho=16.0):
        self.nfig += 1
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(archivo, width=Cm(ancho))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.line_spacing = 1.0
        cap.paragraph_format.space_after = Pt(12)
        fuente(cap.add_run("Figura %d. " % self.nfig), 10, negrita=True)
        fuente(cap.add_run(pie), 10)

    def indice(self):
        self.centrado("Índice", 14, negrita=True, despues=12)
        sec = self.doc.sections[0]
        ancho = sec.page_width - sec.left_margin - sec.right_margin
        for nivel, texto in self.plan_titulos:
            par = self.doc.add_paragraph()
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.left_indent = Cm(0 if nivel == 1 else 0.8)
            par.paragraph_format.tab_stops.add_tab_stop(ancho, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            fuente(par.add_run("%s\t%s" % (texto, self.paginas.get(texto, ""))), 12, negrita=(nivel == 1))

    def numeros_de_pagina(self):
        """Número de página centrado en el pie, sin mostrarlo en la portada."""
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        par = sec.footer.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run()
        for tipo, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if tipo:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), tipo)
            else:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = txt
            run._r.append(e)
        fuente(run, 11)


# ------------------------------------------------------------------ contenido ----

def escribir(inf):
    c, p, h1, h2, lista, tabla, figura = (inf.centrado, inf.p, inf.h1, inf.h2, inf.lista,
                                           inf.tabla, inf.figura)

    # Portada
    c("Pontificia Universidad Católica Madre y Maestra", 16, True, antes=36)
    c("Facultad de Ciencias e Ingeniería", 12, despues=0)
    c("Escuela de Ingeniería en Computación y Telecomunicaciones", 12)
    c("ICC-321-T Inteligencia de Negocios, Grupo 5227", 12, antes=40)
    c("Práctica 1", 18, True, antes=70)
    c("Diseño e implementación de una solución de Inteligencia de Negocios", 15, True, antes=6)
    c("Presentado por:", 12, True, antes=90)
    c("Elías De La Cruz Jiménez, 1015-5971", 12)
    c("Enger Douglas, 1014-4675", 12)
    c("Presentado a:", 12, True, antes=24)
    c("Prof. Lisibonny Beato", 12)
    c("Santiago de los Caballeros, República Dominicana", 12, antes=80)
    c("27 de septiembre de 2026", 12)
    inf.salto()

    inf.indice()
    inf.salto()

    # 1
    h1("1. Introducción")
    p("En esta práctica diseñamos e implementamos una solución de Inteligencia de Negocios completa a "
      "partir de la base de datos operacional BI_Practica_Retail, que se crea con el script entregado junto "
      "al enunciado. La base registra las ventas de una cadena de tiendas al detalle que vende en 12 "
      "sucursales físicas y por dos canales digitales (Web y App), entre el 1 de enero de 2023 y el 15 de "
      "agosto de 2026.")
    p("El trabajo siguió el orden que plantea el enunciado. Primero exploramos la fuente para entender sus "
      "tablas, su granularidad y sus problemas de calidad. Con eso definimos un problema de decisión y las "
      "preguntas de análisis que debía responder la solución. Luego diseñamos un modelo dimensional con la "
      "metodología de Kimball [1], lo implementamos como Data Warehouse en SQL Server, construimos el "
      "proceso ETL en Tableau Prep y, por último, desarrollamos un dashboard interactivo en Tableau Desktop "
      "que lee únicamente del Data Warehouse. Es el recorrido habitual de una solución de BI, que va de la "
      "integración de los datos a su análisis [2]. La Figura 1 lo resume.")
    figura(os.path.join(FIG, "arquitectura.png"), "Arquitectura de la solución, desde la base operacional hasta el dashboard.")
    p("Junto a este documento se entregan el script del Data Warehouse (DW_Retail.sql) y el de carga "
      "(Cargar_DW.sql), el flujo de Tableau Prep empaquetado con sus datos de entrada (ETL_DW_Retail.tflx) y "
      "el libro de Tableau con el dashboard (Dashboard_Retail.twbx), además de las consultas que alimentan "
      "el dashboard (Consultas_Dashboard.sql).")

    # 2
    h1("2. Descripción general de la fuente de datos")
    p("Para crear la base ejecutamos el script de la práctica en SQL Server 2022. El script original no "
      "corrió en el primer intento: en el bloque que genera los detalles de orden usa el alias de columna "
      "LineNo, y LINENO es una palabra reservada de Transact-SQL, por lo que el motor devuelve el error "
      "\"Incorrect syntax near the keyword LineNo\". Renombramos ese alias a LineaNo en sus 9 apariciones y "
      "con ese cambio el script corrió completo. Guardamos esa versión como Retail_Operacional_fix.sql.")
    p("Además notamos que el script no genera siempre los mismos datos. Inserta las filas con claves "
      "IDENTITY desde consultas sin ORDER BY y a partir de esas claves calcula fechas, productos y "
      "devoluciones, así que el resultado depende del orden en que el motor asigne las claves. Cuando lo "
      "ejecutamos en las computadoras de los dos integrantes obtuvimos, por ejemplo, 5,139 devoluciones en "
      "una y 5,181 en la otra. Para que todas las etapas trabajaran sobre los mismos datos, extrajimos las 10 "
      "tablas de una sola ejecución a archivos CSV y todo el trabajo posterior se hizo sobre esa extracción. "
      "Las cifras de este documento corresponden a ella.")

    h2("2.1 Entidades y relaciones")
    p("La base tiene 10 tablas en el esquema dbo, relacionadas por 9 claves foráneas. Seis son tablas de "
      "catálogo, que describen el negocio, y cuatro son transaccionales, que registran lo que ocurre en la "
      "operación (Tabla 1).")
    tabla("Tablas de la base operacional BI_Practica_Retail.",
          ["Tabla", "Tipo", "Filas", "Contenido"], [
              ["Categoria", "Catálogo", "10", "Categorías comerciales de producto"],
              ["Proveedor", "Catálogo", "24", "Proveedores, con país y fecha de alta"],
              ["Producto", "Catálogo", "250", "SKU, nombre, marca, costo unitario, precio de lista y si está activo"],
              ["Tienda", "Catálogo", "12", "Sucursales, con ciudad, provincia, tipo y fecha de apertura"],
              ["Cliente", "Catálogo", "2,500", "Datos demográficos, ubicación, segmento y fecha de registro"],
              ["Promocion", "Catálogo", "15", "Campañas con vigencia, porcentaje de descuento y canal"],
              ["OrdenVenta", "Transaccional", "30,000", "Cabecera de la venta: fecha y hora, tienda, cliente, canal, estado y promoción"],
              ["DetalleOrden", "Transaccional", "76,784", "Líneas de la venta: producto, cantidad, precio unitario y descuento"],
              ["Pago", "Transaccional", "28,803", "Pago de la orden: fecha, método y monto"],
              ["Devolucion", "Transaccional", "5,139", "Devolución de una línea: fecha, cantidad, motivo y reembolso"],
          ], [3.2, 2.8, 1.8, 8.7])
    p("Una orden pertenece a una tienda y, de forma opcional, a un cliente y a una promoción. Cada orden "
      "tiene entre 1 y 4 líneas de detalle y cada línea corresponde a un producto, que a su vez pertenece a "
      "una categoría y a un proveedor. El pago se registra por orden completa, mientras que la devolución se "
      "registra sobre una línea de detalle.")

    h2("2.2 Naturaleza y granularidad de los datos")
    p("Revisamos qué representa una fila de cada tabla transaccional, porque de eso depende el diseño del "
      "modelo. OrdenVenta tiene una fila por venta (30,000 filas y 30,000 identificadores distintos). "
      "DetalleOrden tiene una fila por producto dentro de una orden: 76,784 líneas, con un promedio de 2.56 "
      "líneas por orden; 6,099 órdenes tienen una línea, 8,009 tienen dos, 8,901 tienen tres y 6,991 tienen "
      "cuatro. Pago tiene exactamente un pago por orden completada (28,803 pagos para 28,803 órdenes) y "
      "Devolucion tiene una fila por línea devuelta (5,139 filas con 5,139 líneas distintas).")
    p("La conclusión práctica es que el pago y la devolución no están al mismo grano que la línea de venta. "
      "Si el monto del pago se copiara en cada línea, se repetiría tantas veces como líneas tenga la orden y "
      "cualquier suma quedaría inflada. Por eso los pagos y las devoluciones se trataron como hechos aparte.")

    h2("2.3 Variables disponibles para el análisis")
    p("Las medidas numéricas están en DetalleOrden (cantidad, precio unitario y descuento), en Producto "
      "(costo unitario y precio de lista), en Pago (monto) y en Devolucion (cantidad devuelta y monto "
      "reembolsado). Para analizar esas medidas hay atributos de tiempo (fecha y hora de la orden, fecha de "
      "pago, fecha de devolución), de producto (SKU, nombre, 12 marcas, 10 categorías y 24 proveedores con "
      "su país), de cliente (sexo, fecha de nacimiento, ciudad, 8 provincias y segmento), de tienda (ciudad, "
      "provincia y tipo) y de la venta (canal, estado, método de pago, promoción y motivo de devolución).")
    tabla("Distribuciones observadas en la fuente.",
          ["Variable", "Distribución"], [
              ["Canal de venta", "Tienda 15,223 órdenes (50.7%), Web 8,252 (27.5%) y App 6,525 (21.8%)"],
              ["Estado de la orden", "Completada 28,803 (96.01%) y Cancelada 1,197 (3.99%)"],
              ["Segmento de cliente", "Regular 1,553, Preferente 643 y Corporativo 304 clientes"],
              ["Tipo de tienda", "Calle 4, Centro comercial 4, Express 3 y Outlet 1"],
              ["Promoción", "5,966 órdenes (19.9%) tienen promoción, con descuentos entre 10% y 20%"],
              ["Método de pago", "Cinco métodos, con entre 5,755 y 5,766 pagos cada uno"],
          ], [4.0, 12.5])

    h2("2.4 Problemas de calidad de datos")
    p("Durante la exploración encontramos varias situaciones que había que resolver en el ETL. Ninguna "
      "obligó a descartar datos, salvo las órdenes canceladas (Tabla 3).")
    tabla("Problemas de calidad encontrados y tratamiento aplicado.",
          ["Problema", "Magnitud", "Tratamiento"], [
              ["Órdenes sin cliente (ClienteID nulo)", "2,407 órdenes (8.02%)", "Se asignan a un miembro \"No identificado\" con clave -1 en la dimensión Cliente"],
              ["Clientes sin sexo registrado", "254 de 2,500 (10.2%)", "El nulo se reemplaza por \"No especificado\""],
              ["Órdenes canceladas con líneas de detalle", "1,197 órdenes, 3,117 líneas, RD$12.1 millones", "Se excluyen del hecho de ventas, porque no son ventas efectivas"],
              ["Órdenes sin promoción", "24,034 órdenes (80.1%)", "Se asignan a un miembro \"Sin promoción\" con clave -1"],
              ["Ventas de productos marcados como inactivos", "3,123 líneas de 10 productos", "Se conservan, porque son ventas históricas válidas"],
          ], [5.0, 4.2, 7.3])
    p("También verificamos que no hay devoluciones con más unidades que las vendidas, ni devoluciones o "
      "pagos de órdenes canceladas, que el monto de cada pago coincide con el neto de su orden y que no hay "
      "claves foráneas huérfanas.")

    h2("2.5 Posibilidades de análisis")
    p("Con estos datos se puede analizar la rentabilidad por categoría, marca y proveedor, el efecto de las "
      "promociones, las devoluciones y su costo, el desempeño por canal y tipo de tienda, el comportamiento "
      "de los segmentos de clientes y la evolución de la venta en el tiempo. Como referencia, las órdenes "
      "completadas suman 210,229 unidades, una venta neta de RD$288,991,557.74 y un margen bruto de "
      "RD$75,608,585.04, es decir, 26.16% sobre la venta neta.")

    # 3
    h1("3. Problema de toma de decisiones")
    h2("3.1 Situación encontrada en los datos")
    p("A nivel general el negocio se ve estable: vende entre RD$77 y 82 millones al año y su margen se "
      "mantiene cerca de 26% en todos los años. Sin embargo, cuando ese margen se separa por promoción, por "
      "categoría y por devoluciones aparecen tres problemas que un reporte de ingresos no deja ver.")
    p("El primero es que las promociones reducen mucho el margen sin generar más venta por orden. Las órdenes "
      "con promoción dejan un margen de 17.09%, frente a 28.13% de las que no tienen promoción, y su ticket "
      "promedio es menor (RD$8,965 contra RD$10,299). Las unidades por orden son prácticamente las mismas "
      "(7.29 contra 7.30). El segundo es que la categoría que más vende no es la que más aporta: Papelería "
      "lidera la venta neta con RD$38.7 millones, pero tiene el margen más bajo (23.3%), mientras que Cuidado "
      "personal, con RD$29.2 millones, deja 27.9%. El tercero es que las devoluciones se concentran en pocas "
      "categorías: Moda devuelve el 14.0% de sus líneas y Electrodomésticos el 10.0%, cuando el promedio es "
      "6.98%. En total se reembolsaron RD$13.2 millones, lo que equivale al 17.4% del margen bruto.")

    h2("3.2 Problema y necesidad de información")
    p("La empresa mide su desempeño con información de ingresos, pero no puede medir la rentabilidad real "
      "de sus categorías y productos, porque los descuentos y las devoluciones se registran en procesos "
      "separados de la venta. Por eso puede estar poniendo en promoción categorías cuyo margen no aguanta el "
      "descuento y manteniendo productos con muchas devoluciones sin saber cuánto le cuestan. Lo que hace "
      "falta es una vista integrada donde se puedan seguir juntos la venta, el descuento, el costo, el margen "
      "y los reembolsos, y separarlos por categoría, producto, promoción, canal, tienda, cliente y tiempo.")

    h2("3.3 Usuarios de la solución")
    p("El usuario principal es el gerente comercial o de categorías, que decide el surtido, los precios y "
      "qué categorías entran en promoción. También la usarían el gerente de mercadeo, para evaluar las "
      "campañas; la gerencia de tiendas y canales, para comparar sucursales y el canal físico con el digital; "
      "y la gerencia general, para seguir el margen consolidado.")

    h2("3.4 Decisiones que se busca apoyar")
    lista([
        "Qué categorías, marcas y productos deben entrar o salir de las campañas promocionales.",
        "Qué nivel de descuento puede soportar cada categoría sin perder su margen.",
        "Sobre qué productos actuar para reducir las devoluciones.",
        "Qué proveedores conviene renegociar, según el margen y las devoluciones que generan.",
        "Cómo repartir el esfuerzo comercial entre tiendas y canales.",
    ])

    h2("3.5 Preguntas de análisis")
    p("A partir de lo anterior definimos ocho preguntas, que son las que guían el modelo y el dashboard "
      "(Tabla 4).")
    tabla("Preguntas de análisis.",
          ["#", "Pregunta"], [
              ["P1", "¿Qué categorías, marcas y productos aportan más margen bruto, y no solo más ingresos?"],
              ["P2", "¿Cuánto margen sacrifican las promociones y qué volumen adicional generan realmente?"],
              ["P3", "¿Qué productos y categorías concentran las devoluciones y cuánto cuestan en reembolsos?"],
              ["P4", "¿Cuál es el margen neto de devoluciones por categoría y por proveedor?"],
              ["P5", "¿Cómo evolucionan la venta y el margen mes a mes, y qué meses concentran la demanda?"],
              ["P6", "¿Qué canal de venta y qué tipo de tienda rinden mejor en ticket promedio y margen?"],
              ["P7", "¿Qué segmentos de cliente sostienen la venta y qué peso tienen las ventas sin cliente identificado?"],
              ["P8", "¿Qué proveedores concentran el margen y cuáles arrastran las devoluciones?"],
          ], [1.3, 15.2])

    # 4
    h1("4. Diseño del modelo dimensional")
    h2("4.1 Proceso de negocio y grano")
    p("El proceso de negocio que modelamos es la venta al detalle de productos en tiendas y canales "
      "digitales. Siguiendo a Kimball y Ross [1], el grano se declara antes de elegir medidas y dimensiones. "
      "El grano de la tabla de hechos principal es una línea de detalle de una orden de venta completada, es "
      "decir, una fila por cada producto vendido dentro de una orden. Elegimos este grano porque es el nivel "
      "más fino que ofrece la fuente y permite resumir hacia arriba (por orden, producto, categoría, tienda, "
      "día o mes) sin perder información. Las 1,197 órdenes canceladas y sus 3,117 líneas quedan fuera porque "
      "no representan ventas.")

    h2("4.2 Medidas")
    p("Las medidas de FactVentas se obtienen de la línea de detalle y del costo del producto (Tabla 5). Seis "
      "de ellas son aditivas y se pueden sumar por cualquier dimensión. El precio unitario y el margen "
      "porcentual no son aditivos: el primero se promedia y el segundo se calcula en el dashboard a partir de "
      "sus componentes.")
    tabla("Medidas de FactVentas.",
          ["Medida", "Cómo se obtiene", "Aditividad"], [
              ["CantidadVendida", "DetalleOrden.Cantidad", "Aditiva"],
              ["MontoBruto", "Cantidad × PrecioUnitario", "Aditiva"],
              ["DescuentoMonto", "DetalleOrden.DescuentoMonto", "Aditiva"],
              ["MontoNeto", "MontoBruto - DescuentoMonto", "Aditiva"],
              ["CostoTotal", "Cantidad × Producto.CostoUnitario", "Aditiva"],
              ["MargenBruto", "MontoNeto - CostoTotal", "Aditiva"],
              ["PrecioUnitario", "DetalleOrden.PrecioUnitario", "No aditiva (se promedia)"],
              ["MargenPorcentual", "MargenBruto / MontoNeto", "No aditiva (no se almacena)"],
          ], [4.0, 7.0, 5.5])

    h2("4.3 Dimensiones")
    tabla("Dimensiones del modelo.",
          ["Dimensión", "Atributos principales", "Observaciones"], [
              ["DimFecha", "Fecha, año, trimestre, mes, semana, día de la semana, fin de semana", "Cubre del 1-1-2023 al 31-12-2026. Se usa como fecha de venta, de pago y de devolución"],
              ["DimProducto", "SKU, nombre, marca, categoría, proveedor, país del proveedor, rango de precio, activo", "Categoría y proveedor se incluyen dentro de la dimensión (esquema estrella, no copo de nieve)"],
              ["DimCliente", "Nombre, sexo, grupo de edad, ciudad, provincia, segmento, año de registro", "Incluye el miembro \"No identificado\" con clave -1"],
              ["DimTienda", "Nombre, ciudad, provincia, tipo de tienda, año de apertura", "12 sucursales"],
              ["DimPromocion", "Nombre, porcentaje de descuento, canal, vigencia", "Incluye el miembro \"Sin promoción\" con clave -1"],
              ["DimCanal", "Canal de venta", "Tienda, Web y App"],
              ["DimMotivoDevolucion", "Motivo", "Cinco motivos. Solo aplica a las devoluciones"],
              ["DimMetodoPago", "Método de pago", "Cinco métodos. Solo aplica a los pagos"],
          ], [4.3, 6.4, 5.8])

    h2("4.4 Esquema estrella")
    p("La Figura 2 muestra el esquema estrella de FactVentas con sus seis dimensiones. OrdenID se guarda en "
      "el hecho como dimensión degenerada, porque identifica la orden pero no tiene atributos propios que "
      "justifiquen una tabla aparte.")
    figura(os.path.join(FIG, "estrella.png"), "Esquema estrella de FactVentas.", 12.0)
    p("El modelo se completa con dos hechos más, que no se pueden meter en FactVentas porque están a otro "
      "grano: FactDevoluciones, con una fila por devolución de una línea, y FactPagos, con una fila por pago "
      "de una orden. Los tres hechos comparten las dimensiones de fecha, cliente y tienda, lo que permite "
      "filtrarlos de la misma forma. La Tabla 7 es la matriz de bus del modelo.")
    tabla("Matriz de bus.",
          ["Hecho", "Fecha", "Producto", "Cliente", "Tienda", "Promoción", "Canal", "Motivo", "Método de pago"], [
              ["FactVentas", "X", "X", "X", "X", "X", "X", "", ""],
              ["FactDevoluciones", "X", "X", "X", "X", "", "", "X", ""],
              ["FactPagos", "X", "", "X", "X", "", "X", "", "X"],
          ], [3.3, 1.4, 1.8, 1.6, 1.5, 2.2, 1.4, 1.6, 1.7])

    h2("4.5 Claves y relaciones")
    p("Cada dimensión tiene una clave subrogada entera, propia del Data Warehouse. En este caso la "
      "asignamos a partir del identificador del origen y además guardamos la clave natural como atributo, "
      "para poder rastrear de dónde viene cada fila. DimFecha usa como clave un entero con formato "
      "AAAAMMDD, que se lee fácil y es la práctica que recomienda Kimball [1]. Los hechos guardan solo las "
      "claves foráneas, las medidas y la dimensión degenerada, y todas las relaciones son de uno a muchos "
      "desde la dimensión hacia el hecho. Los miembros con clave -1 evitan que queden hechos sin relación "
      "cuando el origen trae nulos.")

    h2("4.6 Relación entre preguntas, medidas y dimensiones")
    tabla("Trazabilidad entre preguntas, medidas y dimensiones.",
          ["Pregunta", "Medidas", "Dimensiones"], [
              ["P1", "MargenBruto, MontoNeto, margen %", "Producto (categoría, marca), Fecha"],
              ["P2", "MontoNeto, DescuentoMonto, MargenBruto, CantidadVendida", "Promoción, Producto, Fecha"],
              ["P3", "CantidadDevuelta, MontoReembolso", "Producto, Motivo, Fecha"],
              ["P4", "MargenBruto - MontoReembolso", "Producto (categoría, proveedor)"],
              ["P5", "MontoNeto, MargenBruto", "Fecha (año, trimestre, mes)"],
              ["P6", "MontoNeto, MargenBruto, ticket promedio", "Canal, Tienda (tipo de tienda)"],
              ["P7", "MontoNeto, CantidadVendida", "Cliente (segmento, provincia)"],
              ["P8", "MargenBruto, MontoReembolso", "Producto (proveedor, país)"],
          ], [2.2, 7.2, 7.1])

    # 5
    h1("5. Implementación del Data Warehouse")
    p("El modelo se implementó en SQL Server con el script DW_Retail.sql, que crea la base DW_Retail desde "
      "cero. Las ocho dimensiones tienen su clave subrogada como PRIMARY KEY y los tres hechos tienen una "
      "clave propia de tipo bigint tomada del identificador de la línea, la devolución o el pago del origen. "
      "En total hay 16 claves foráneas (6 en FactVentas, 5 en FactDevoluciones y 5 en FactPagos), de manera "
      "que ningún hecho puede apuntar a un miembro que no exista.")
    p("Para los montos usamos decimal(14,2) y para los precios decimal(12,2), así no se pierden centavos. "
      "Todas las columnas de los hechos son NOT NULL, y las cantidades vendidas y devueltas tienen un CHECK "
      "que exige que sean mayores que cero. También creamos 8 índices sobre las claves foráneas que más se "
      "usan en las consultas (fecha, producto, cliente, tienda y promoción).")
    p("DimFecha se llena en el mismo script con los 1,461 días del 1 de enero de 2023 al 31 de diciembre de "
      "2026, sin depender de los hechos. Antes de llenarla fijamos SET DATEFIRST 7, porque el número del día "
      "de la semana y de la semana del año dependen de esa configuración del servidor [4]. El script también "
      "inserta los miembros \"No identificado\" y \"Sin promoción\" con clave -1.")
    p("La carga la hace Cargar_DW.sql, que lee los CSV que produce el flujo de Tableau Prep con BULK INSERT "
      "[3]. Como Tableau Prep escribe las columnas en orden alfabético, los archivos se cargan primero en "
      "tablas temporales y de ahí se insertan en las definitivas con el tipo de dato correcto. Durante esta "
      "parte tuvimos que resolver algunos problemas que dependían de la computadora donde se ejecutaba:")
    lista([
        "En SQL Server para Linux la opción CODEPAGE de BULK INSERT no está disponible, así que la quitamos "
        "después de comprobar que los datos no tienen caracteres fuera de ASCII.",
        "Tableau Prep escribe las fechas como mes/día/año. Para que se lean igual en cualquier servidor, el "
        "script fija SET DATEFORMAT mdy.",
        "Tableau Prep en Windows termina las líneas de los CSV con CRLF y en macOS con LF. Después de cada "
        "carga se elimina el retorno de carro que queda en la última columna, y así el script funciona con "
        "archivos de cualquiera de las dos computadoras.",
        "La carpeta de los CSV se indica en la variable RutaCSV al principio del script, que se ejecuta en "
        "modo SQLCMD.",
    ])
    p("Después de la carga comparamos el Data Warehouse con la fuente. Los conteos y los totales coinciden "
      "(Tabla 9). La prueba más importante es la de la venta neta: FactVentas y FactPagos suman exactamente "
      "RD$288,991,557.74, lo mismo que las órdenes completadas del origen, lo que confirma que no se perdió "
      "ni se duplicó ninguna línea al pasar de un grano al otro.")
    tabla("Comprobación de la carga del Data Warehouse.",
          ["Tabla", "Filas", "Comprobación"], [
              ["FactVentas", "73,667", "76,784 líneas menos las 3,117 de órdenes canceladas"],
              ["FactDevoluciones", "5,139", "Todas las devoluciones del origen"],
              ["FactPagos", "28,803", "Un pago por cada orden completada"],
              ["DimCliente", "2,501", "2,500 clientes más el miembro \"No identificado\""],
              ["DimPromocion", "16", "15 promociones más el miembro \"Sin promoción\""],
              ["DimProducto / DimTienda", "250 / 12", "Iguales al catálogo de origen"],
              ["DimCanal / DimMotivoDevolucion / DimMetodoPago", "3 / 5 / 5", "Valores distintos del origen"],
              ["DimFecha", "1,461", "Todos los días de 2023 a 2026"],
          ], [5.8, 2.4, 8.3])

    # 6
    h1("6. Proceso ETL en Tableau Prep")
    p("El ETL se construyó en Tableau Prep Builder 2026.2 (Figura 3). El flujo tiene 40 nodos: 10 entradas, "
      "una por cada tabla extraída del origen; 11 pasos de limpieza; 6 uniones; 3 agregaciones y 10 salidas, "
      "una por cada tabla del Data Warehouse. Las salidas se escriben como CSV en la carpeta CSV_salida y de "
      "ahí las carga Cargar_DW.sql. El flujo se entrega empaquetado (.tflx), de modo que al abrirlo incluye "
      "los archivos de entrada [7].")
    p("El flujo se puede leer en tres bloques. Arriba están las dimensiones de catálogo: DimProducto une "
      "Producto con Categoria y con Proveedor, y DimCliente, DimTienda y DimPromocion limpian cada tabla por "
      "separado. En el medio están las dimensiones que salen de los datos transaccionales: DimCanal, "
      "DimMotivoDevolucion y DimMetodoPago se obtienen agregando los valores distintos de las órdenes, las "
      "devoluciones y los pagos. Abajo están los hechos, que parten de las órdenes completadas y se unen con "
      "el detalle, el producto, las devoluciones y los pagos [6]. La Tabla 10 resume las transformaciones y "
      "la razón de cada una.")
    figura(os.path.join(REPO, "5_ETL", "capturas", "Flujo_Tableau_Prep.png"),
           "Flujo ETL_DW_Retail en Tableau Prep Builder.", 11.0)
    tabla("Transformaciones del flujo ETL.",
          ["Paso", "Transformación", "Motivo"], [
              ["Producto + Categoria + Proveedor", "Dos uniones internas por CategoriaID y ProveedorID", "DimProducto lleva la categoría y el proveedor dentro (esquema estrella)"],
              ["DimProducto", "Clave subrogada y rango de precio (Económico hasta RD$800, Medio hasta RD$1,500, Premium)", "Permite agrupar el margen por rango de precio"],
              ["DimCliente", "Une nombre y apellido, reemplaza el sexo nulo, calcula la edad al 15 de agosto de 2026 y el grupo de edad, y el año de registro", "Atributos para analizar a los clientes (P7). La fecha fija hace que el grupo no cambie según el día en que se corre el flujo"],
              ["DimTienda, DimPromocion", "Clave subrogada y año de apertura", "Dimensiones de P2 y P6"],
              ["DimCanal, DimMotivoDevolucion, DimMetodoPago", "Agregación por valor distinto y clave fija para cada valor", "En el origen son solo texto dentro de las transacciones"],
              ["Filtrar órdenes completadas", "Filtro EstadoOrden = \"Completada\", clave de fecha AAAAMMDD, cliente y promoción nulos como -1, clave de canal", "Grano del hecho y miembros -1"],
              ["Detalle + Orden + Producto", "Une cada línea con su orden y con el costo unitario del producto", "El costo está en el catálogo, no en la línea"],
              ["FactVentas", "Calcula MontoBruto, MontoNeto, CostoTotal y MargenBruto redondeados a dos decimales", "Las medidas quedan listas para sumarse"],
              ["FactDevoluciones", "Une la devolución con su línea, clave de fecha de la devolución y clave del motivo", "Hecho aparte, al grano de la devolución"],
              ["FactPagos", "Une el pago con su orden, clave de fecha del pago y clave del método", "Hecho al grano de la orden"],
          ], [4.0, 6.5, 6.0])
    p("Un detalle que nos tomó tiempo fue el del redondeo. Al principio Tableau Prep escribía algunos "
      "márgenes en notación científica (por ejemplo 2.27e-13) cuando el resultado era prácticamente cero, y "
      "SQL Server no podía convertir ese texto. Lo resolvimos aplicando ROUND(..., 2) a las medidas "
      "monetarias dentro del flujo, que es además lo correcto para importes.")
    p("Para validar el ETL comparamos las salidas con el origen línea por línea. Los 73,667 registros de "
      "FactVentas tienen la misma venta neta, el mismo cliente, la misma promoción y la misma fecha que su "
      "línea de origen, y ningún registro quedó sin clave de canal, motivo o método de pago.")

    # 7
    h1("7. Análisis de las medidas")
    p("En esta sección explicamos cada medida: qué representa, cómo se obtiene, su relación con el grano, "
      "por qué dimensiones se puede analizar, qué preguntas responde y cómo se puede agregar. Los totales son "
      "de todo el período cargado.")
    h2("7.1 Medidas de FactVentas")
    p("Todas las medidas de FactVentas están al grano de la línea de orden completada: cada fila las aporta "
      "una sola vez, así que cualquier resumen hacia arriba es una suma directa. Se pueden analizar por las "
      "seis dimensiones del hecho (fecha, producto con su categoría y proveedor, cliente, tienda, promoción "
      "y canal) y responden las preguntas P1, P2, P5, P6 y P7.")
    tabla("Medidas de FactVentas con su total del período.",
          ["Medida", "Qué representa y cómo se obtiene", "Total", "Agregación"], [
              ["CantidadVendida", "Unidades del producto en la línea (DetalleOrden.Cantidad)", "210,229 unidades", "Aditiva"],
              ["MontoBruto", "Valor de la línea antes del descuento: Cantidad × PrecioUnitario", "RD$298,749,966.36", "Aditiva"],
              ["DescuentoMonto", "Descuento aplicado a la línea (DetalleOrden.DescuentoMonto)", "RD$9,758,408.62", "Aditiva"],
              ["MontoNeto", "Lo que realmente se cobra: MontoBruto - DescuentoMonto", "RD$288,991,557.74", "Aditiva"],
              ["CostoTotal", "Costo de la mercancía: Cantidad × CostoUnitario del producto", "RD$213,382,972.70", "Aditiva"],
              ["MargenBruto", "Ganancia antes de devoluciones: MontoNeto - CostoTotal", "RD$75,608,585.04", "Aditiva"],
              ["PrecioUnitario", "Precio cobrado por unidad en la línea", "Promedio simple RD$1,421.91", "No aditiva"],
          ], [3.3, 6.6, 3.6, 3.0])
    p("El precio unitario no se puede sumar. Además, su promedio simple (RD$1,421.91) no es igual al precio "
      "promedio real ponderado por unidades, que es MontoBruto entre CantidadVendida (RD$1,421.07). Por eso "
      "en el dashboard los precios y los porcentajes se calculan a partir de las medidas aditivas y no se "
      "promedian valores ya calculados.")
    p("También hay que tener en cuenta que DescuentoMonto no viene solo de las promociones. Las ventas sin "
      "promoción tienen algunos descuentos puntuales del 5%: son 7,159 líneas que suman RD$1.5 millones, el "
      "0.63% de su venta bruta. Es poco, por lo que el grupo \"Sin promoción\" sigue sirviendo como punto de "
      "comparación en la pregunta P2.")

    h2("7.2 Medidas de FactDevoluciones")
    p("CantidadDevuelta (9,543 unidades en total) y MontoReembolso (RD$13,158,252.77) se copian de la tabla "
      "Devolucion y están al grano de la devolución de una línea. Ambas son aditivas por fecha, producto, "
      "cliente, tienda y motivo, y responden P3, P4 y P8. La fecha de este hecho es la de la devolución, no "
      "la de la venta, por eso hay devoluciones hasta el 3 de septiembre de 2026 aunque la última venta es del "
      "15 de agosto. Estas medidas no se pueden cruzar con promoción ni con canal, porque ese hecho no tiene "
      "esas dimensiones.")

    h2("7.3 Medida de FactPagos")
    p("MontoPagado (RD$288,991,557.74) está al grano de la orden, porque en el negocio se paga la orden "
      "completa y no cada línea. Es aditiva por fecha, cliente, tienda, canal y método de pago, pero no se "
      "puede repartir por producto ni por categoría. Su total es igual a la venta neta, lo que nos sirvió "
      "como control de la carga. La fecha es la del pago: en 1,727 de las 28,803 órdenes el pago ocurrió otro "
      "día, y en 56 de ellas en otro mes, así que al comparar por día o por mes los totales de FactPagos y "
      "FactVentas pueden no coincidir aunque el total del período sí coincida.")

    h2("7.4 Indicadores calculados en el dashboard")
    p("Los indicadores que son razones no se guardan en el Data Warehouse; se calculan en Tableau con sus "
      "componentes, al nivel de detalle que se esté viendo (Tabla 12).")
    tabla("Indicadores calculados a partir de las medidas.",
          ["Indicador", "Fórmula", "Valor del período", "Preguntas"], [
              ["Margen %", "SUM(MargenBruto) / SUM(MontoNeto)", "26.16%", "P1, P2, P6, P8"],
              ["Órdenes", "COUNTD(OrdenID)", "28,803", "P6"],
              ["Ticket promedio", "SUM(MontoNeto) / COUNTD(OrdenID)", "RD$10,033", "P2, P6"],
              ["Tasa de devolución", "Devoluciones / líneas vendidas", "6.98%", "P3, P8"],
              ["Margen después de devoluciones", "SUM(MargenBruto) - SUM(MontoReembolso)", "RD$62,450,332.27", "P4"],
              ["Margen que se va en reembolsos", "SUM(MontoReembolso) / SUM(MargenBruto)", "17.40%", "P4, P8"],
          ], [4.3, 6.2, 3.4, 2.6])
    p("Las órdenes se cuentan con COUNTD y nunca se suman entre grupos, porque una misma orden tiene varias "
      "líneas. Para el margen después de devoluciones no se puede unir FactVentas con FactDevoluciones fila "
      "por fila, porque las líneas con devolución se duplicarían. Lo que hicimos fue resumir cada hecho por "
      "separado al mismo nivel (año y producto) y luego combinarlos por esas dimensiones compartidas, que es "
      "lo que Kimball llama consulta entre hechos o drill-across [1].")

    # 8
    h1("8. Dashboard interactivo")
    p("El dashboard se hizo en Tableau Desktop y se entrega como libro empaquetado (Dashboard_Retail.twbx). "
      "Sus datos salen únicamente de DW_Retail: las consultas de Consultas_Dashboard.sql unen cada hecho con "
      "sus dimensiones y el resultado se guardó como un extracto de Tableau que va dentro del libro [5]. Así "
      "el libro se puede abrir en cualquier computadora con Tableau sin conectarse al servidor, y ningún dato "
      "viene de la base operacional.")
    p("Cada tablero tiene filtros de año y categoría, y el primero también de canal. Los hicimos con "
      "parámetros de Tableau [8] porque un parámetro filtra al mismo tiempo las dos fuentes de datos del "
      "libro, y los aplicamos como filtro de contexto para que el ranking de productos se calcule dentro de lo "
      "filtrado. Al pasar el cursor por cualquier barra o punto se ve el detalle del valor.")

    h2("8.1 Tablero 1: rentabilidad de la venta")
    figura(os.path.join(REPO, "7_Dashboard", "capturas", "Tablero_1_Rentabilidad.png"),
           "Tablero \"Rentabilidad de la venta\", sin filtros.", 16.0)
    p("Las cifras de arriba resumen el período: venta neta de RD$289.0 millones, margen bruto de RD$75.6 "
      "millones (26.2%), 28,803 órdenes y un ticket promedio de RD$10,033. Para P1, el gráfico de categorías "
      "muestra que Papelería es la que más vende (RD$38.7 millones) pero la de menor margen (23.3%), mientras "
      "que Cuidado personal (27.9%) y Electrodomésticos (27.4%) son las más rentables.")
    p("Para P2, el margen cae a medida que la campaña da más descuento: San Valentín (10%) deja 20.4%, Madres "
      "(12%) 18.9%, Regreso a clases (15%) 15.9% y Black Friday (20%) apenas 11.0%, frente a 28.1% sin "
      "promoción. Como el ticket con promoción es menor, el descuento no se está compensando con más venta.")
    p("Para P5, la venta mensual se mueve entre RD$5.4 y 7.7 millones, con un promedio de RD$6.65 millones, "
      "y los meses más altos fueron marzo de 2025 y julio de 2024 y de 2025. Donde sí se ve un patrón es en el "
      "margen: agosto de cada año baja a cerca de 16% (15.8% en 2023 y 16.1% en 2024 y 2025), cuando corre la "
      "campaña de Regreso a clases, y noviembre de 2025, con Black Friday, baja a 22.7%. En el gráfico esos "
      "meses aparecen en un tono más claro. La última barra es baja porque agosto de 2026 solo tiene datos "
      "hasta el día 15. Para P6, el ticket promedio va de RD$9,385 a RD$10,205 y el margen de 25.8% a 26.5% en todas las "
      "combinaciones de canal y tipo de tienda; la mayor parte de la venta ocurre en tiendas físicas de centro "
      "comercial. Para P7, el segmento Regular aporta el 56.3% de la venta, Preferente el 24.3% y Corporativo "
      "el 11.5%, y el 7.8% de la venta no tiene cliente identificado.")

    h2("8.2 Tablero 2: devoluciones y proveedores")
    figura(os.path.join(REPO, "7_Dashboard", "capturas", "Tablero_2_Devoluciones.png"),
           "Tablero \"Devoluciones y proveedores\", sin filtros.", 16.0)
    p("En este tablero se ve que los reembolsos suman RD$13.2 millones, con una tasa de devolución de 7.0%, "
      "y que después de restarlos el margen baja de RD$75.6 a 62.5 millones: el 17.4% del margen se va en "
      "reembolsos. Para P3, Moda (14.0%), Electrodomésticos (10.0%) y Tecnología (9.3%) tienen tasas que "
      "duplican a las demás categorías, que están entre 4.7% y 5.1%, y los diez productos con más dinero "
      "reembolsado son de esas mismas categorías; el primero es el Producto 188, de Moda, con RD$232,410.")
    p("Para P4, Moda pierde el 35.9% de su margen en reembolsos, Electrodomésticos el 25.0% y Tecnología el "
      "24.6%, mientras que Papelería, a pesar de tener el margen más bajo, solo pierde el 14.5%. Para P8, el "
      "gráfico de dispersión muestra que el Proveedor 03 tiene uno de los márgenes más altos (29.7%) pero "
      "también la mayor tasa de devolución (11.2%), que el Proveedor 13 tiene el margen más bajo (18.8%) y que "
      "los proveedores 24 y 01 combinan márgenes por encima de 31% con devoluciones bajas.")
    figura(os.path.join(REPO, "7_Dashboard", "capturas", "Tablero_2_Filtro_Moda.png"),
           "El mismo tablero con el filtro de categoría en Moda.", 16.0)
    p("La Figura 6 muestra la interacción. Al elegir Moda en el filtro de categoría, todas las vistas se "
      "recalculan: la tasa de devolución de la categoría (14.0%), el margen que pierde (35.9%), los diez "
      "productos de Moda con más reembolsos y los proveedores que la abastecen.")

    # 9
    h1("9. Justificación de las principales decisiones de diseño")
    p("Estas son las decisiones que más influyeron en la solución y por qué las tomamos:")
    lista([
        "**Grano en la línea de orden completada.** Si el hecho estuviera a nivel de orden se perdería el "
        "producto, que es el eje de P1, P3 y P4.",
        "**Tres hechos separados.** El pago es por orden y la devolución es por línea; meterlos en la misma "
        "tabla que la venta duplicaría montos y rompería la aditividad.",
        "**Excluir las órdenes canceladas.** No son ventas efectivas; incluirlas sumaría RD$12.1 millones que "
        "nunca se cobraron.",
        "**Miembros con clave -1.** Si el cliente o la promoción quedaran en nulo, 2,407 órdenes no se podrían "
        "cruzar con la dimensión Cliente y no se podría comparar con y sin promoción.",
        "**Esquema estrella y no copo de nieve.** Tener la categoría y el proveedor dentro de DimProducto "
        "reduce las uniones y hace el modelo más fácil de usar.",
        "**Medidas derivadas calculadas en el ETL.** Quedan listas en el Data Warehouse para cualquier "
        "herramienta, y el dashboard solo tiene que sumarlas.",
        "**Razones calculadas en el dashboard.** Un porcentaje guardado no se puede sumar ni promediar bien "
        "cuando se agrupa.",
        "**Filtros con parámetros.** Un parámetro filtra las dos fuentes del dashboard a la vez y permite que "
        "el ranking de productos se calcule dentro del filtro.",
    ])

    # 10
    h1("10. Conclusiones")
    p("La solución permite ver en un solo lugar lo que antes estaba separado entre ventas, descuentos y "
      "devoluciones. Con eso se responde el problema planteado: la empresa no puede juzgar sus categorías y "
      "promociones solo por lo que venden.")
    p("Los resultados más importantes para la toma de decisiones son tres. Las promociones profundas no se "
      "justifican con estos datos, porque Black Friday deja 11.0% de margen y no sube el ticket, así que "
      "conviene limitar los descuentos grandes a las categorías con más margen. Papelería necesita una "
      "revisión de precios o de costos con sus proveedores, porque es la categoría que más vende con el menor "
      "margen. Y las devoluciones de Moda, Electrodomésticos y Tecnología son la mayor fuente de pérdida de "
      "margen, por lo que los productos del ranking del segundo tablero son el punto de partida para revisar "
      "la calidad y la información de los productos, junto con proveedores como el 03. En cambio, el canal y "
      "el tipo de tienda casi no cambian la rentabilidad, y la caída del margen cada agosto confirma que el "
      "efecto de las promociones se ve incluso en la evolución mensual.")
    p("En cuanto al proceso, lo que más aprendimos fue la importancia de definir bien el grano desde el "
      "principio: casi todas las decisiones posteriores, desde separar los hechos hasta la forma de calcular "
      "el margen después de devoluciones, salieron de esa definición. También quedó claro que hay que "
      "verificar las cifras en cada etapa, porque problemas como el script que cambia en cada ejecución o las "
      "diferencias de formato entre computadoras solo aparecieron al comparar los resultados.")

    # Referencias
    inf.titulos.append((1, "Referencias"))
    inf.doc.add_heading("Referencias", level=1)
    refs = [
        "R. Kimball y M. Ross, *The Data Warehouse Toolkit: The Definitive Guide to Dimensional Modeling*, "
        "3.a ed. Indianapolis, IN, EE. UU.: Wiley, 2013.",
        "R. Sherman, *Business Intelligence Guidebook: From Data Integration to Analytics*. Waltham, MA, "
        "EE. UU.: Morgan Kaufmann, 2014.",
        "Microsoft, \"BULK INSERT (Transact-SQL),\" Microsoft Learn. [En línea]. Disponible: "
        "https://learn.microsoft.com/en-us/sql/t-sql/statements/bulk-insert-transact-sql",
        "Microsoft, \"SET DATEFIRST (Transact-SQL),\" Microsoft Learn. [En línea]. Disponible: "
        "https://learn.microsoft.com/en-us/sql/t-sql/statements/set-datefirst-transact-sql",
        "Tableau, \"Extract Your Data,\" Tableau Desktop Help. [En línea]. Disponible: "
        "https://help.tableau.com/current/pro/desktop/en-us/extracting_data.htm",
        "Tableau, \"Aggregate, Join, or Union Data,\" Tableau Prep Help. [En línea]. Disponible: "
        "https://help.tableau.com/current/prep/en-us/prep_combine.htm",
        "Tableau, \"Save and Share Tableau Prep Work,\" Tableau Prep Help. [En línea]. Disponible: "
        "https://help.tableau.com/current/prep/en-us/prep_save_share.htm",
        "Tableau, \"Create Parameters,\" Tableau Desktop Help. [En línea]. Disponible: "
        "https://help.tableau.com/current/pro/desktop/en-us/parameters_create.htm",
    ]
    for i, r in enumerate(refs, 1):
        par = inf.doc.add_paragraph()
        par.paragraph_format.left_indent = Cm(1.0)
        par.paragraph_format.first_line_indent = Cm(-1.0)
        par.paragraph_format.line_spacing = 1.15
        par.paragraph_format.space_after = Pt(6)
        texto_con_formato(par, "[%d]\t%s" % (i, r))


# ------------------------------------------------------------------ armado ----

PLAN = None


def construir(paginas):
    inf = Informe(paginas)
    inf.plan_titulos = PLAN or []
    escribir(inf)
    inf.numeros_de_pagina()
    cp = inf.doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.comments = ""
    cp.title = "Práctica 1: Diseño e implementación de una solución de Inteligencia de Negocios"
    cp.subject = "ICC-321-T Inteligencia de Negocios"
    cp.keywords = ""
    cp.language = "es-DO"
    return inf


def a_libreoffice(docx_entrada, carpeta):
    """Pasa el .docx por LibreOffice (como un documento guardado desde Writer) y saca el PDF."""
    tmp = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--convert-to", "odt", "--outdir", tmp, docx_entrada],
                   check=True, capture_output=True)
    odt = os.path.join(tmp, os.path.splitext(os.path.basename(docx_entrada))[0] + ".odt")
    subprocess.run(["soffice", "--headless", "--convert-to", "docx:MS Word 2007 XML", "--outdir", carpeta, odt],
                   check=True, capture_output=True)
    final = os.path.join(carpeta, os.path.basename(docx_entrada))
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", carpeta, final],
                   check=True, capture_output=True)
    shutil.rmtree(tmp)
    return final, os.path.splitext(final)[0] + ".pdf"


def paginas_de_titulos(pdf, titulos):
    """Busca en qué página del PDF aparece cada título (ignorando portada e índice)."""
    salida = {}
    n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout).group(1))
    textos = [subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), "-layout", pdf, "-"],
                             capture_output=True, text=True).stdout for i in range(1, n + 1)]
    for _, t in titulos:
        for i in range(2, n):
            lineas = [l.strip() for l in textos[i].splitlines()]
            if t in lineas:
                salida[t] = i + 1
                break
    return salida


def main():
    global PLAN
    tmp = tempfile.mkdtemp()
    borrador = os.path.join(tmp, NOMBRE + ".docx")
    inf = construir({})
    PLAN = inf.titulos
    inf = construir({})
    inf.doc.save(borrador)
    _, pdf = a_libreoffice(borrador, tmp)
    paginas = paginas_de_titulos(pdf, PLAN)
    faltan = [t for _, t in PLAN if t not in paginas]
    if faltan:
        print("Títulos sin página:", faltan)
    inf = construir(paginas)
    inf.doc.save(borrador)
    docx_final, pdf_final = a_libreoffice(borrador, SALIDA)
    print(docx_final)
    print(pdf_final)
    shutil.rmtree(tmp)


if __name__ == "__main__":
    main()
